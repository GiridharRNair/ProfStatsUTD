#!/usr/bin/env python3
"""Import UTD grade distributions into Supabase.

The parser can run without credentials using --dry-run. Uploads require:

    SUPABASE_URL
    SUPABASE_SECRET_KEY
"""

from __future__ import annotations

import argparse
import csv
import os
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterator


GRADE_COLUMNS = {
    "a_plus": ("A+",),
    "a": ("A",),
    "a_minus": ("A-",),
    "b_plus": ("B+",),
    "b": ("B",),
    "b_minus": ("B-",),
    "c_plus": ("C+",),
    "c": ("C",),
    "c_minus": ("C-",),
    "d_plus": ("D+",),
    "d": ("D",),
    "d_minus": ("D-",),
    "f": ("F",),
    "cr": ("CR",),
    "nc": ("NC",),
    "p": ("P",),
    "w": ("W", "Total W", "W Total"),
    "i": ("I",),
    "nf": ("NF",),
}

EDGE_CASE_PROFESSORS = {
    "Du, Ding Z": "Ding-Zhu Du",
    "Nishi, Hirofumi": "Hiro Nishi",
    "Macalevey, Paul J": "Paul MacAlevey",
    "Mac Alevey, Paul J": "Paul MacAlevey",
    "Razo-Razo, Miguel Angel": "Miguel Razo",
}

REQUIRED_COLUMNS = ("Subject", "Section", "Instructor 1")
COURSE_NUMBER_COLUMNS = ("Catalog Nbr", "Catalog Number")
SUPPORTED_SUFFIXES = {".csv", ".xlsx"}


@dataclass(frozen=True)
class GradeSection:
    term: str
    subject: str
    catalog_number: str
    section: str
    instructor_name: str
    instructor_search_name: str
    grades: dict[str, int]
    source_file: str

    def to_supabase_record(self) -> dict[str, object]:
        return {
            "term": self.term,
            "subject": self.subject,
            "catalog_number": self.catalog_number,
            "section": self.section,
            "instructor_name": self.instructor_name,
            "instructor_search_name": self.instructor_search_name,
            "source_file": self.source_file,
            **self.grades,
        }


def reorder_name(name: str) -> str:
    name = name.strip()
    if name in EDGE_CASE_PROFESSORS:
        return EDGE_CASE_PROFESSORS[name]

    name_parts = [part.strip() for part in name.split(",", maxsplit=1)]
    if len(name_parts) == 1:
        return name_parts[0]

    last_name, rest = name_parts
    return f"{rest} {last_name}".strip()


def search_name(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", name.lower()).strip()


def normalize_header(header: object) -> str:
    return str(header or "").strip().lstrip("\ufeff")


def normalize_row(row: dict[object, object]) -> dict[str, str]:
    return {normalize_header(key): str(value or "").strip() for key, value in row.items()}


def first_value(row: dict[str, str], *columns: str) -> str:
    for column in columns:
        value = row.get(column, "").strip()
        if value:
            return value
    return ""


def int_value(row: dict[str, str], *columns: str) -> int:
    value = first_value(row, *columns)
    if not value:
        return 0

    return int(float(value))


def term_from_path(path: Path) -> str:
    return path.stem.strip()


def validate_columns(row: dict[str, str], source_file: str) -> None:
    missing_columns = [column for column in REQUIRED_COLUMNS if column not in row]
    if not any(column in row for column in COURSE_NUMBER_COLUMNS):
        missing_columns.append("Catalog Nbr or Catalog Number")

    if missing_columns:
        raise ValueError(f"{source_file} is missing required column(s): {', '.join(missing_columns)}")


def parse_grade_row(row: dict[str, str], source_path: Path) -> GradeSection:
    validate_columns(row, source_path.name)

    instructor_name = reorder_name(row["Instructor 1"])
    grades = {grade_key: int_value(row, *columns) for grade_key, columns in GRADE_COLUMNS.items()}

    return GradeSection(
        term=term_from_path(source_path),
        subject=row["Subject"].upper(),
        catalog_number=first_value(row, *COURSE_NUMBER_COLUMNS).upper(),
        section=row["Section"].upper(),
        instructor_name=instructor_name,
        instructor_search_name=search_name(instructor_name),
        grades=grades,
        source_file=source_path.name,
    )


def read_csv(path: Path) -> Iterator[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as csv_file:
        for row in csv.DictReader(csv_file):
            yield normalize_row(row)


def read_xlsx(path: Path) -> Iterator[dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Install openpyxl to import .xlsx files: python3 -m pip install openpyxl") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    headers = [normalize_header(header) for header in next(rows)]

    for values in rows:
        yield normalize_row(dict(zip(headers, values)))


def read_source(path: Path) -> Iterator[dict[str, str]]:
    if path.suffix.lower() == ".csv":
        yield from read_csv(path)
        return

    if path.suffix.lower() == ".xlsx":
        yield from read_xlsx(path)
        return

    raise ValueError(f"Unsupported grade file type: {path}")


def iter_source_files(data_dir: Path) -> Iterator[Path]:
    for path in sorted(data_dir.iterdir()):
        if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES:
            yield path


def parse_sources(data_dir: Path) -> list[GradeSection]:
    sections: list[GradeSection] = []
    for source_path in iter_source_files(data_dir):
        for row in read_source(source_path):
            section = parse_grade_row(row, source_path)
            sections.append(section)

    return sections


def chunked(records: list[dict[str, object]], size: int) -> Iterator[list[dict[str, object]]]:
    for index in range(0, len(records), size):
        yield records[index : index + size]


def create_supabase_client():
    supabase_url = os.environ.get("SUPABASE_URL")
    secret_key = os.environ.get("SUPABASE_SECRET_KEY")

    if not supabase_url or not secret_key:
        raise RuntimeError("Set SUPABASE_URL and SUPABASE_SECRET_KEY before uploading.")

    try:
        from supabase import create_client
    except ImportError as exc:
        raise RuntimeError("Install supabase before uploading: python3 -m pip install supabase") from exc

    return create_client(supabase_url, secret_key)


def upload_records(records: list[dict[str, object]], batch_size: int) -> None:
    client = create_supabase_client()
    conflict_key = "source_file,subject,catalog_number,section,instructor_name"

    for batch in chunked(records, batch_size):
        client.table("grade_sections").upsert(batch, on_conflict=conflict_key).execute()


def summarize(sections: list[GradeSection]) -> dict[str, int]:
    professors = {section.instructor_search_name for section in sections}
    courses = {(section.subject, section.catalog_number) for section in sections}
    terms = {section.term for section in sections}

    return {
        "rows": len(sections),
        "professors": len(professors),
        "courses": len(courses),
        "terms": len(terms),
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Import UTD grade distribution files into Supabase.")
    parser.add_argument("--data-dir", default="raw_data", type=Path, help="Directory containing .csv or .xlsx grade files.")
    parser.add_argument("--dry-run", action="store_true", help="Parse files and print a summary without uploading.")
    parser.add_argument("--batch-size", default=500, type=int, help="Supabase upsert batch size.")
    return parser


def main() -> None:
    args = build_parser().parse_args()
    sections = parse_sources(args.data_dir)
    summary = summarize(sections)

    print(
        "Parsed {rows} rows across {professors} professors, {courses} courses, and {terms} terms.".format(
            **summary
        )
    )

    if args.dry_run:
        return

    records = [section.to_supabase_record() for section in sections]
    upload_records(records, args.batch_size)
    print(f"Uploaded {len(records)} rows to Supabase.")


if __name__ == "__main__":
    main()
